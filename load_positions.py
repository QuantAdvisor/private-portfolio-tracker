"""
Phase 0 – Positions-Loader
Liest ETF Portfolio/Portfolio.xlsx und schreibt einen datierten Snapshot
in portfolio.position_snapshot.

Aufruf:
    python load_positions.py [--date YYYY-MM-DD]

Ohne --date wird das heutige Datum verwendet.
Idempotent: mehrfacher Lauf am selben Tag überschreibt den Snapshot.
"""

import argparse
import logging
import sys
from datetime import date
from pathlib import Path

import openpyxl
import pandas as pd

import db_utils

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(message)s",
    datefmt="%H:%M:%S",
)
log = logging.getLogger(__name__)

PORTFOLIO_FILE = Path(__file__).parent / "ETF Portfolio" / "Portfolio.xlsx"
SHEET = "Tabelle1"


# ── Einlesen ──────────────────────────────────────────────────────────────────

def load_portfolio_xlsx() -> pd.DataFrame:
    wb = openpyxl.load_workbook(PORTFOLIO_FILE, data_only=True)
    ws = wb[SHEET]
    rows = [row for row in ws.iter_rows(values_only=True) if any(v is not None for v in row)]

    header = rows[0]
    data   = rows[1:]
    df = pd.DataFrame(data, columns=header)

    # Spaltennamen normalisieren (Tippfehler in Quelldatei: 'Isisn', 'Sparrate im Monata')
    df = df.rename(columns={
        "Isisn":                  "isin",
        "Depot":                  "depot",
        "Kategorie":              "kategorie",
        "Name":                   "name",
        "Stück":                  "quantity",
        "Avg_Kaufkurs":           "avg_cost",
        "Sparrate im Monata":     "sparrate",
        "akuteller kurs (EUR)":   "price_eur",
    })

    # Nur Zeilen mit echten Positionen behalten
    df = df[df["isin"].notna() & df["depot"].notna()].copy()

    # ISIN säubern: führendes \xa0 (non-breaking space) entfernen
    df["isin"] = df["isin"].astype(str).str.strip()

    # avg_cost: 'unbekannt' (Riester) → None
    df["avg_cost"] = pd.to_numeric(df["avg_cost"], errors="coerce")

    # quantity sicherstellen
    df["quantity"] = pd.to_numeric(df["quantity"], errors="coerce")

    missing_qty = df[df["quantity"].isna()]
    if not missing_qty.empty:
        log.warning("Stückzahl fehlt für %d Zeilen – werden übersprungen:", len(missing_qty))
        for _, r in missing_qty.iterrows():
            log.warning("  %s | %s", r["isin"], r["name"])
        df = df[df["quantity"].notna()]

    return df


# ── Depot-ID nachschlagen ─────────────────────────────────────────────────────

def get_account_map() -> dict[str, int]:
    rows = db_utils.query_df("SELECT account_id, name FROM portfolio.account")
    return dict(zip(rows["name"], rows["account_id"]))


# ── ETF-ISIN-Existenz prüfen ──────────────────────────────────────────────────

def check_etf_exists(isins: list[str]) -> set[str]:
    placeholders = ", ".join(f":isin_{i}" for i in range(len(isins)))
    params = {f"isin_{i}": v for i, v in enumerate(isins)}
    result = db_utils.query_df(
        f"SELECT isin FROM portfolio.etf WHERE isin IN ({placeholders})", params
    )
    return set(result["isin"].tolist())


# ── Snapshot schreiben ────────────────────────────────────────────────────────

UPSERT_SQL = """
INSERT INTO portfolio.position_snapshot
    (account_id, isin, as_of_date, quantity, avg_cost)
VALUES
    (:account_id, :isin, :as_of_date, :quantity, :avg_cost)
ON CONFLICT (account_id, isin, as_of_date)
DO UPDATE SET
    quantity  = EXCLUDED.quantity,
    avg_cost  = EXCLUDED.avg_cost,
    loaded_at = now()
"""

def write_snapshot(rows: list[dict]) -> None:
    db_utils.execute_many(UPSERT_SQL, rows)


# ── Hauptprogramm ─────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="Lade Portfolio.xlsx in portfolio.position_snapshot")
    parser.add_argument(
        "--date",
        type=date.fromisoformat,
        default=date.today(),
        help="Snapshot-Datum (YYYY-MM-DD), Standard: heute",
    )
    args = parser.parse_args()
    as_of = args.date
    log.info("Snapshot-Datum: %s", as_of)

    # 1. Datei einlesen
    log.info("Lese %s ...", PORTFOLIO_FILE)
    df = load_portfolio_xlsx()
    log.info("%d Positionen gelesen", len(df))

    # 2. Depot-Mapping
    account_map = get_account_map()
    unknown_depots = set(df["depot"]) - set(account_map)
    if unknown_depots:
        log.error("Unbekannte Depots in Portfolio.xlsx: %s", unknown_depots)
        log.error("Bitte in portfolio.account nachtragen und erneut laufen.")
        sys.exit(1)

    # 3. ETF-ISINs prüfen
    all_isins = df["isin"].unique().tolist()
    known_isins = check_etf_exists(all_isins)
    missing_isins = set(all_isins) - known_isins
    if missing_isins:
        log.warning(
            "%d ISINs nicht in portfolio.etf – werden übersprungen "
            "(bitte in db_schema.sql nachtragen):",
            len(missing_isins),
        )
        for isin in sorted(missing_isins):
            name = df.loc[df["isin"] == isin, "name"].iloc[0]
            log.warning("  %s  %s", isin, name)
        df = df[df["isin"].isin(known_isins)]

    if df.empty:
        log.error("Keine ladbaren Positionen übrig. Abbruch.")
        sys.exit(1)

    # 4. Snapshot-Zeilen bauen
    records = []
    for _, row in df.iterrows():
        records.append({
            "account_id": account_map[row["depot"]],
            "isin":       row["isin"],
            "as_of_date": as_of,
            "quantity":   float(row["quantity"]),
            "avg_cost":   float(row["avg_cost"]) if pd.notna(row["avg_cost"]) else None,
        })

    # 5. In DB schreiben
    log.info("Schreibe %d Zeilen → portfolio.position_snapshot ...", len(records))
    try:
        write_snapshot(records)
    except Exception as exc:
        log.exception("Fehler beim DB-Schreiben: %s", exc)
        sys.exit(1)

    log.info("Fertig. %d Positionen für %s gespeichert.", len(records), as_of)

    # 6. Kurze Zusammenfassung
    summary = (
        df.groupby("depot")
        .agg(etfs=("isin", "count"), gesamtstück=("quantity", "sum"))
        .reset_index()
    )
    log.info("\n%s", summary.to_string(index=False))


if __name__ == "__main__":
    main()
